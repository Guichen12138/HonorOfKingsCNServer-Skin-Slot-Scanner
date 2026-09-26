import requests
import concurrent.futures
import threading
import time
import json
import random
import os
from email.utils import parsedate_to_datetime
from datetime import timezone, timedelta

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# 共用常量与工具：Excel 路径 / 槽位范围 / 表头解析 / URL 拼装
# （common.py 和本脚本在同一目录，Python 会自动找到）
from common import (
    EXCEL_FILE,
    SKIN_START,
    SKIN_END,
    parse_header_columns,
    skin_image_url,
)


# ============================================================
# 配置
# ============================================================

# 英雄编号范围
HERO_START = 1
HERO_END = 999

# 并发线程
#
# 2026-09-12 实测（64/128/192/256/320 线程基准）：
#   64→440 req/s、128→770、192→800、256→791、320→832
# 吞吐在 ~192 线程后到顶（CDN 单 IP 并发上限），
# 192 是性价比点，全量扫描约 25 秒。
# （此前"96 无增益"的结论已过时，CDN 后来放行了更多并发）
THREADS = 192

# 请求超时时间
TIMEOUT = 10


# 单个地址失败后的重试次数
RETRIES = 3

# 每次重试前的等待秒数（按重试次数递增）
RETRY_DELAY = 0.5

# --------------------------------------------------------------------
# 缓存标记
#
# 每一轮扫描使用同一个值，让它不重复、但整轮保持一致。
#
# 2026-09-26 起只在重试时带上它：
#
# 之前所有请求都带 ?_=时间戳，等于每次扫描都全量回源，
# CDN 忙的时候（中午）一轮要 18~20 秒，深夜只要 8~10 秒。
# 现在首次请求不带标记，让边缘节点正常命中（快且稳定）；
# 只有重试（404 确认 / 429 / 5xx / 无 Last-Modified）才带上，
# 强制回源拿最新结果。
#
# 保留这个标记的初衷不变，关键点在于「过期的 404」：
# 资源还没上线时扫出来的是 404，边缘节点会把这个 404 缓存下来；
# 等资源真正上传之后，边缘节点可能还在给旧的 404，
# 于是这个最新的资源就一直扫不到，Last-Modified 也就一直写不进去。
# 对应的处理见 check_skin 里的 404 分支。
# --------------------------------------------------------------------

CACHE_BUSTER = str(int(time.time() * 1000))


# ============================================================
# Session
# ============================================================

thread_local = threading.local()


def get_session():

    if not hasattr(thread_local, "session"):

        session = requests.Session()

        session.headers.update({
            "User-Agent": "Mozilla/5.0",

            # 让 CDN 回源校验，
            # 不要直接返回边缘节点上的过期副本
            "Cache-Control": "no-cache",
            "Pragma": "no-cache"
        })

        thread_local.session = session

    return thread_local.session


# ============================================================
# 全局 429 冷却（重试风暴的刹车）
#
# 之前每个线程撞到 429 后各睡各的（固定递增延时），
# 192 个线程几乎同时睡醒、同时重试，
# 刚缓过来的 CDN 立刻又被压垮，形成重试风暴。
#
# 现在任何线程撞到 429，就全局推后一个冷却窗口；
# 所有线程发请求前先看一眼，窗口没过就先等。
# ============================================================

_throttle_lock = threading.Lock()

# 冷却截止时刻（time.monotonic() 秒）
_throttle_until = 0.0


def _note_429():

    global _throttle_until

    with _throttle_lock:

        _throttle_until = max(
            _throttle_until,
            time.monotonic() + 1.0
        )


def _wait_throttle():

    while True:

        with _throttle_lock:

            wait = (
                _throttle_until
                - time.monotonic()
            )

        if wait <= 0:

            return

        time.sleep(
            min(wait, 0.5)
        )


# ============================================================
# 判断一个 404 响应是不是可能来自边缘节点的旧缓存
#
# 腾讯 CDN 的标记头是 X-Cache-Lookup，而且常常多段连在一起：
#   "Cache Hit, Hit From Inner Cluster, Cache Miss"
# 只要里面出现过 Cache Miss，说明这条链路上刚回源过，
# 这个 404 就是源头当前的真实答案，不用再确认；
# 只有全是 Hit 时才可能是「过期的 404」，需要带 buster 回源确认。
#
# 兼容其他 CDN 的标准 X-Cache / Age 头；
# 什么线索都没有时保守当成缓存，宁可多确认一次。
# ============================================================

def _may_be_stale_cache(response):

    lookup = (
        response.headers.get("X-Cache-Lookup")
        or ""
    )

    if lookup:

        return "MISS" not in lookup.upper()

    x_cache = (
        response.headers.get("X-Cache")
        or ""
    ).upper()

    if x_cache:

        if "MISS" in x_cache:

            return False

        if "HIT" in x_cache:

            return True

    age = response.headers.get("Age")

    try:

        return int(age) > 0

    except (TypeError, ValueError):

        return True


# ============================================================
# 检查单个资源
# ============================================================

def check_skin(hero_id, skin_id):

    url = skin_image_url(hero_id, skin_id)


    # ================================================================
    # 为什么改成 GET（和 curl -i 一样），不再用 HEAD
    #
    # HEAD 在这个 CDN 上命中的缓存和 GET 不是同一份，
    # 刚更新过的资源经常在 HEAD 上拿到旧的缓存结果，
    # 于是最新的资源扫不出来，Last-Modified 也就一直写不进去。
    #
    # 用 GET 时加 stream=True：
    # 只读取响应头，不下载图片正文，开销和 HEAD 差不多。
    # ================================================================

    for attempt in range(
        1,
        RETRIES + 1
    ):

        # 有线程撞到 429 时全局正在冷却，
        # 先等一等再发，别火上浇油
        _wait_throttle()

        try:

            session = get_session()

            response = session.get(
                url,
                timeout=TIMEOUT,
                allow_redirects=True,
                stream=True,

                # 首次请求不带缓存标记，
                # 让边缘节点正常命中（快且不受 CDN 忙时影响）；
                # 重试才带上，强制回源拿最新结果
                params=(
                    {"_": CACHE_BUSTER}
                    if attempt > 1
                    else None
                ),
            )

            try:

                status = response.status_code


                # ------------------------------------------------
                # 429 / 5xx 是服务器临时繁忙或限流，
                # 不能当成「资源不存在」，要重试
                #
                # 404 之类的才是真的不存在
                # ------------------------------------------------

                if (
                        status == 429
                        or 500 <= status < 600
                ):

                    if attempt < RETRIES:

                        # 读掉小错误体：
                        # 不读的话这个连接会被整个丢弃，
                        # 下一个请求又要重新 TCP+TLS 握手，
                        # 这正是之前扫描慢的主要原因
                        response.content

                        # 429 记入全局冷却窗口，
                        # 让所有线程都缓一缓再发
                        if status == 429:
                            _note_429()

                        # 抖动：把 192 个线程的重试时刻错开，
                        # 不要同时睡醒、同时打回去
                        time.sleep(
                            RETRY_DELAY * attempt
                            + random.uniform(0.1, 0.6)
                        )

                        continue


                    # 重试次数用完还是失败：
                    # 标记为失败，让统计能看出来

                    response.content

                    return {
                        "hero_id": hero_id,
                        "skin_id": skin_id,
                        "exists": False,
                        "failed": True
                    }


                if status != 200:

                    # 同上：读掉 404 的小错误体
                    # （几百字节的 XML，很便宜），
                    # 让连接可以复用给下一个请求
                    response.content

                    # ------------------------------------------------
                    # 「过期的 404」确认（2026-09-26 起只在 404 时做）：
                    #
                    # 首次 404 且响应可能来自边缘缓存时，
                    # 带 buster 强制回源再确认一次；
                    # X-Cache: MISS 说明刚回源，404 可信，
                    # 空号英雄不用白打第二个请求。
                    # ------------------------------------------------

                    if (
                            status == 404
                            and attempt == 1
                            and _may_be_stale_cache(response)
                    ):
                        continue

                    return {
                        "hero_id": hero_id,
                        "skin_id": skin_id,
                        "exists": False
                    }

                last_modified = response.headers.get(
                    "Last-Modified",
                    ""
                )


                # ------------------------------------------------
                # 200 但没有 Last-Modified：
                #
                # 说明这次拿到的响应不完整，
                # 不能当成「存在但没有时间」写进表，
                # 直接重试
                # ------------------------------------------------

                if not last_modified:

                    time.sleep(
                        RETRY_DELAY * attempt
                        + random.uniform(0.1, 0.6)
                    )

                    continue


                try:
                    dt = parsedate_to_datetime(last_modified)

                    # GMT/UTC → 北京时间 UTC+8
                    dt = dt.astimezone(
                        timezone(timedelta(hours=8))
                    )

                    last_modified = dt.strftime(
                        "%Y-%m-%d %H:%M:%S"
                    )

                except Exception:
                    pass


                return {
                    "hero_id": hero_id,
                    "skin_id": skin_id,
                    "exists": True,
                    "last_modified": last_modified
                }


            finally:

                # 只取了响应头，
                # 这里必须关掉，否则连接不会释放
                response.close()


        except requests.RequestException:

            # 偶发抖动：等一下再试，
            # 不要因为一次失败就把这条记录整个丢掉
            if attempt < RETRIES:

                time.sleep(
                    RETRY_DELAY * attempt
                    + random.uniform(0.1, 0.6)
                )

                continue


    # 重试全部用完还是失败：
    # 标记为失败，让统计能看出来，
    # 提醒用户这一格本次没扫到

    return {
        "hero_id": hero_id,
        "skin_id": skin_id,
        "exists": False,
        "failed": True
    }


# ============================================================
# 扫描一个英雄
# ============================================================

def scan_hero(hero_id):

    results = []

    # 这一英雄有多少个格子重试完还是失败
    failed = 0


    # ================================================================
    # 先探测 00 原皮槽，空号英雄直接跳过剩余槽位
    #
    # 依据：任何英雄上线必然带原皮（00 槽），
    # 现有 139 个英雄的起始槽全部是 00，无一例外。
    # 00 槽 404 → 这个英雄编号不存在 → 后面 19 个请求全是白打。
    #
    # 2026-09-14 改造后请求量：
    #   19980 → 约 3640（存在英雄×20 + 空号×1），降幅 82%。
    #
    # 注意：00 槽「请求失败」（429/5xx/超时重试用完）
    # 不能当成空号，会保守地继续扫完这个英雄，
    # 防止把真实英雄误判成不存在。
    # ================================================================

    first = check_skin(
        hero_id,
        SKIN_START
    )

    if first["exists"]:

        results.append(first)

    elif first.get("failed"):

        # 00 槽请求失败：不断定英雄不存在，
        # 继续扫完整个英雄，按原逻辑统计失败数

        failed += 1

    else:

        # 00 槽 404：英雄编号不存在，跳过剩余槽位

        return hero_id, results, failed


    for skin_id in range(
        SKIN_START + 1,
        SKIN_END + 1
    ):

        result = check_skin(
            hero_id,
            skin_id
        )

        if result["exists"]:

            results.append(result)

        elif result.get("failed"):

            failed += 1

    return hero_id, results, failed


# ============================================================
# 官网英雄列表（英雄名自动填充用）
#
# ename 就是本工具用的英雄编号，cname 是官方英雄名。
# 编号体系已用全表手输名字比对验证（2026-09-25，95%+ 一致；
# 少数不一致全是手输名的笔误/注记，本功能只填空格，不受影响）。
# ============================================================

HERO_LIST_URL = "https://pvp.qq.com/web201605/js/herolist.json"


def fetch_official_hero_names():

    """
    拉官网英雄列表，返回 {英雄编号: 官方英雄名}。

    任何失败都返回空 dict：
    自动填充是锦上添花，不能影响扫描主流程。
    """

    try:

        response = requests.get(
            HERO_LIST_URL,
            timeout=TIMEOUT,
            headers={"User-Agent": "Mozilla/5.0"},
        )

        response.raise_for_status()

        content = response.content

        try:
            data = json.loads(content.decode("utf-8"))

        except UnicodeDecodeError:
            data = json.loads(content.decode("gbk"))

        names = {}

        for item in data:

            try:

                names[
                    int(item["ename"])
                ] = str(item["cname"]).strip()

            except (KeyError, TypeError, ValueError):

                continue

        return names

    except Exception as exc:

        print(
            f"官网英雄列表获取失败，"
            f"本次跳过英雄名自动填充：{exc}"
        )

        return {}


# ============================================================
# 主程序
# ============================================================

def main():

    print("=" * 70)
    print("王者荣耀 HeroOriginalPainting 增量扫描器")
    print("=" * 70)
    print()


    # ========================================================
    # 打开 / 创建 Excel
    # ========================================================

    if os.path.exists(EXCEL_FILE):

        excel_existed = True

        print(
            f"发现已有 Excel：{EXCEL_FILE}"
        )

        print(
            "正在读取已有数据……"
        )

        wb = load_workbook(
            EXCEL_FILE
        )

        ws = wb.active

        print(
            "已有数据将被保留，不会初始化覆盖。"
        )

    else:

        excel_existed = False

        print(
            "没有找到 Excel，正在创建新的文件……"
        )

        wb = Workbook()

        ws = wb.active

        ws.title = "皮肤资源"


    # ========================================================
    # 确保基本表头存在
    # ========================================================

    ws.cell(
        row=1,
        column=1,
        value="英雄编号"
    )

    ws.cell(
        row=1,
        column=2,
        value="英雄名字"
    )


    # ========================================================
    # 获取已有列
    # ========================================================

    skin_name_columns, last_modified_columns = parse_header_columns(ws)


    # ========================================================
    # 增加缺失的皮肤格子列
    # ========================================================

    for skin_id in range(
        SKIN_START,
        SKIN_END + 1
    ):

        # ----------------------------------------------------
        # 如果没有这个皮肤名列，就新增
        # ----------------------------------------------------

        if skin_id not in skin_name_columns:

            new_col = ws.max_column + 1

            ws.cell(
                row=1,
                column=new_col,
                value=f"{skin_id:02d}-皮肤名"
            )

            skin_name_columns[
                skin_id
            ] = new_col


        # ----------------------------------------------------
        # 如果没有 Last-Modified 列，就新增
        # ----------------------------------------------------

        if skin_id not in last_modified_columns:

            new_col = ws.max_column + 1

            ws.cell(
                row=1,
                column=new_col,
                value=f"{skin_id:02d}-Last-Modified"
            )

            last_modified_columns[
                skin_id
            ] = new_col


    # ========================================================
    # 找已有英雄
    # ========================================================

    hero_rows = {}


    for row in range(
        2,
        ws.max_row + 1
    ):

        hero_id = ws.cell(
            row=row,
            column=1
        ).value

        if hero_id is None:
            continue

        try:

            hero_id = int(hero_id)

            hero_rows[
                hero_id
            ] = row

        except ValueError:

            pass


    print()

    print(
        f"Excel 中已有 "
        f"{len(hero_rows)} 个英雄。"
    )

    print()

    print(
        f"本次扫描英雄："
        f"{HERO_START:03d} ~ {HERO_END:03d}"
    )

    print(
        f"本次扫描皮肤槽："
        f"{SKIN_START:02d} ~ {SKIN_END:02d}"
    )

    print()

    # ========================================================
    # 开始扫描
    # ========================================================

    total_heroes = (
        HERO_END -
        HERO_START +
        1
    )

    all_results = {}

    # 重试完还是失败的请求数
    total_failed = 0

    # 整个英雄任务崩掉的数量
    scan_errors = 0

    # 00 槽 404、被直接跳过的空号英雄数量
    skipped_heroes = 0

    # 是否被 Ctrl+C 手动中断
    interrupted = False


    executor = concurrent.futures.ThreadPoolExecutor(
        max_workers=THREADS
    )

    futures = {
        executor.submit(
            scan_hero,
            hero_id
        ): hero_id
        for hero_id in range(
            HERO_START,
            HERO_END + 1
        )
    }


    completed = 0


    try:

        for future in concurrent.futures.as_completed(
            futures
        ):

            # ------------------------------------------------
            # 单个英雄的任务崩了，
            # 记一笔继续跑，
            # 不能让整轮扫描的结果全部丢掉
            # ------------------------------------------------

            try:

                hero_id, results, failed = future.result()

            except Exception:

                scan_errors += 1

                completed += 1

                continue

            completed += 1

            total_failed += failed


            if results:

                all_results[
                    hero_id
                ] = results

            elif not failed:

                # 00 槽 404、无任何失败：空号英雄，被跳过

                skipped_heroes += 1


            print(
                f"\r扫描进度："
                f"{completed}/{total_heroes}",
                end="",
                flush=True
            )


    except KeyboardInterrupt:

        # ------------------------------------------------
        # Ctrl+C：取消还没开始的任务，
        # 已扫完的部分照常保存，不白跑
        # ------------------------------------------------

        interrupted = True

        executor.shutdown(
            wait=False,
            cancel_futures=True
        )


    else:

        executor.shutdown(wait=True)


    print()
    print()

    if interrupted:

        print(
            f"已手动中断：本次只扫完 "
            f"{completed}/{total_heroes} 个英雄。"
        )

        print(
            "已扫完的部分会照常保存，"
            "建议之后完整重跑一次。"
        )

    else:

        print(
            "扫描完成，开始更新 Excel……"
        )

    print()


    # ========================================================
    # 更新 Excel
    # ========================================================

    # 和其他已有数据格子一样的填充色
    cell_fill = PatternFill(
        fill_type="solid",
        fgColor="E2F0D9"
    )

    new_heroes = 0

    # 原来 没有 → 这次有：新增
    new_records = 0

    # 原来 有 → 这次时间变了：更新
    updated_records = 0


    for hero_id in sorted(
        all_results
    ):

        results = all_results[
            hero_id
        ]


        # ----------------------------------------------------
        # 判断英雄是否已经存在
        # ----------------------------------------------------

        if hero_id in hero_rows:

            row = hero_rows[
                hero_id
            ]

        else:

            # 新英雄 → 新增一行

            row = ws.max_row + 1

            ws.cell(
                row=row,
                column=1,
                value=hero_id
            )

            # 英雄名字故意留空
            # 给你自己填写

            ws.cell(
                row=row,
                column=2,
                value=""
            )

            hero_rows[
                hero_id
            ] = row

            new_heroes += 1


        # ----------------------------------------------------
        # 更新每一个实际存在的皮肤
        # ----------------------------------------------------

        for result in results:

            skin_id = result[
                "skin_id"
            ]

            last_modified = result[
                "last_modified"
            ]


            name_col = skin_name_columns[
                skin_id
            ]

            time_col = last_modified_columns[
                skin_id
            ]


            # ------------------------------------------------
            # 皮肤名：
            #
            # 绝对不覆盖已有内容
            # ------------------------------------------------

            skin_name = ws.cell(
                row=row,
                column=name_col
            ).value


            # 如果原来就是空的，就保持空
            #
            # 如果已经填写，就完全不碰
            #
            # 所以这里什么都不用做

            old_time = ws.cell(
                row=row,
                column=time_col
            ).value

            # ============================================================
            # Last-Modified：
            # 每次都以服务器最新返回值为准
            #
            # 除了值不同要更新外，
            # 如果旧格子不是字符串类型（比如在 Excel/WPS 里被重新
            # 输入过、变成了真日期单元格），也强制重写回标准字符串，
            # 让显示格式和其他格子保持一致。
            # ============================================================

            if last_modified:

                if (
                        old_time is None
                        or not isinstance(old_time, str)
                        or str(old_time).strip() != str(last_modified).strip()
                ):
                    time_cell = ws.cell(
                        row=row,
                        column=time_col,
                        value=last_modified
                    )

                    # 文本格式（@）：
                    # 即使之后在 Excel/WPS 里重新输入这个格子，
                    # 也保持文本，不会被自动转成日期格式

                    time_cell.number_format = "@"

                    name_cell = ws.cell(
                        row=row,
                        column=name_col
                    )

                    # 出现新的 Last-Modified：
                    # 给 编号-Last-Modified 和 编号-皮肤名 两个格子
                    # 涂上和其他格子一样的颜色
                    time_cell.fill = cell_fill
                    name_cell.fill = cell_fill

                    # 区分「新增」和「更新」两种情况

                    if old_time is None:

                        new_records += 1

                    else:

                        updated_records += 1


    # ========================================================
    # 英雄名自动填充（官网英雄列表）
    #
    # 只填空格，已有名字一个都不碰；
    # 只填扫到过资源的英雄行，
    # 编号还没上线的空行不填。
    # ========================================================

    auto_named = 0

    official_names = fetch_official_hero_names()

    # 00 槽的 Last-Modified 列：
    # 有值说明这个编号真实存在
    first_time_col = (
        last_modified_columns.get(SKIN_START)
        if official_names
        else None
    )

    if official_names and first_time_col is not None:

        for row in range(
            2,
            ws.max_row + 1
        ):

            hero_id = ws.cell(
                row=row,
                column=1
            ).value

            if hero_id is None:
                continue

            try:

                hero_id = int(hero_id)

            except (TypeError, ValueError):

                continue

            # 没扫到过任何资源的行：
            # 编号未上线的占位行，不填
            if not ws.cell(
                row=row,
                column=first_time_col
            ).value:
                continue

            name_cell = ws.cell(
                row=row,
                column=2
            )

            # 已有名字一律不碰
            if (
                name_cell.value is not None
                and str(name_cell.value).strip() != ""
            ):
                continue

            official_name = official_names.get(
                hero_id
            )

            if official_name:

                name_cell.value = official_name

                # 和机器写入的其他格子一样的填充色，
                # 方便一眼看出哪些名字是自动填的
                name_cell.fill = cell_fill

                auto_named += 1

    # ========================================================
    # 样式
    # ========================================================

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7"
    )

    header_font = Font(
        bold=True
    )


    for cell in ws[1]:

        cell.fill = header_fill

        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


    # ========================================================
    # 设置列宽
    # ========================================================

    ws.column_dimensions[
        "A"
    ].width = 12

    ws.column_dimensions[
        "B"
    ].width = 15


    for skin_id in range(
        SKIN_START,
        SKIN_END + 1
    ):

        name_col = skin_name_columns[
            skin_id
        ]

        time_col = last_modified_columns[
            skin_id
        ]


        ws.column_dimensions[
            get_column_letter(
                name_col
            )
        ].width = 18


        ws.column_dimensions[
            get_column_letter(
                time_col
            )
        ].width = 22


    # ========================================================
    # 冻结
    # ========================================================

    ws.freeze_panes = "C2"


    # ========================================================
    # 保存
    #
    # 数据没有变化时跳过保存：
    # openpyxl 每次 save 都会重新打包整个 xlsx，
    # 即使内容一字不差，文件字节也会变，
    # git 就会一直误报「Excel 有未提交的修改」。
    # 只有真正有新增/更新，或文件是本次新建的，才落盘。
    # ========================================================

    has_changes = (
        new_heroes
        or new_records
        or updated_records
        or auto_named
    )

    if has_changes or not excel_existed:

        wb.save(
            EXCEL_FILE
        )

    else:

        print(
            "数据无变化，跳过写入 Excel，"
            "git 不会再误报文件改动。"
        )

        print()


    # ========================================================
    # 统计
    # ========================================================

    print("=" * 70)

    print(
        "更新完成！"
    )

    print("=" * 70)

    print()

    print(
        f"新增英雄："
        f"{new_heroes}"
    )

    print(
        f"自动填充英雄名："
        f"{auto_named}"
    )

    print(
        f"跳过空号英雄："
        f"{skipped_heroes}"
    )

    print(
        f"新增 Last-Modified："
        f"{new_records}"
    )

    print(
        f"更新 Last-Modified："
        f"{updated_records}"
    )

    print()

    if total_failed or scan_errors:

        print(
            f"注意：{total_failed} 个请求、"
            f"{scan_errors} 个英雄任务失败，"
            f"这些格子本次没有扫到。"
        )

        print(
            "建议再完整跑一次补齐。"
        )

        print()

    print(
        f"Excel：{EXCEL_FILE}"
    )

    print()

    print(
        "已有皮肤名不会被修改。"
    )

    print(
        "已有 Last-Modified 以服务器最新值为准，"
        "有更新会自动刷新并高亮。"
    )

    print()


if __name__ == "__main__":

    main()