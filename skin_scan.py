import openpyxl
import webbrowser
import os

# 共用常量与工具：Excel 路径 / 槽位范围 / 表头解析 / URL 拼装
# （common.py 和本脚本在同一目录，Python 会自动找到）
from common import (
    EXCEL_FILE,
    SKIN_START,
    SKIN_END,
    parse_header_columns,
    skin_image_url,
    pause,
)


# ============================================================
# 检查 Excel
# ============================================================

if not os.path.exists(EXCEL_FILE):

    print(f"找不到 Excel 文件：{EXCEL_FILE}")

    pause()
    exit()


# ============================================================
# 读取 Excel
# ============================================================

wb = openpyxl.load_workbook(
    EXCEL_FILE,
    data_only=True
)

ws = wb.active


# ============================================================
# 找到各个 Last-Modified 和皮肤名列
# ============================================================

skin_name_columns, last_modified_columns = parse_header_columns(ws)


# ============================================================
# 读取英雄
# ============================================================

heroes = []

# 实际有效的英雄数量
# （不直接用 ws.max_row - 1，
#   那样会把中间或末尾的空行也算进去）
total_hero_count = 0


for row in range(
    2,
    ws.max_row + 1
):

    hero_id = ws.cell(
        row=row,
        column=1
    ).value

    hero_name = ws.cell(
        row=row,
        column=2
    ).value


    if hero_id is None:
        continue


    try:

        hero_id = int(
            hero_id
        )

    except ValueError:

        continue

    total_hero_count += 1


    # ========================================================
    # 检查这个英雄有哪些实际存在的资源
    # ========================================================

    skins_to_open = []


    for skin_id in range(
        SKIN_START,
        SKIN_END + 1
    ):

        last_modified_col = (
            last_modified_columns.get(
                skin_id
            )
        )

        skin_name_col = (
            skin_name_columns.get(
                skin_id
            )
        )


        if (
            last_modified_col is None
            or skin_name_col is None
        ):
            continue


        # ----------------------------------------------------
        # 服务器资源是否存在
        # ----------------------------------------------------

        last_modified = ws.cell(
            row=row,
            column=last_modified_col
        ).value


        if not last_modified:
            continue


        # ----------------------------------------------------
        # 你是否已经填写过皮肤名字
        # ----------------------------------------------------

        skin_name = ws.cell(
            row=row,
            column=skin_name_col
        ).value


        # 有名字 → 已经处理过 → 跳过
        if (
            skin_name is not None
            and str(skin_name).strip() != ""
        ):
            continue


        # ----------------------------------------------------
        # 有资源 + 没填写名字
        # → 需要打开
        # ----------------------------------------------------

        url = skin_image_url(hero_id, skin_id)


        skins_to_open.append({
            "skin_id": skin_id,
            "url": url,
            "last_modified": last_modified
        })


    # ========================================================
    # 只有存在“未处理资源”的英雄才加入列表
    # ========================================================

    if skins_to_open:

        heroes.append({
            "row": row,
            "hero_id": hero_id,
            "hero_name": (
                hero_name
                if hero_name
                else "未填写英雄名"
            ),
            "skins": skins_to_open
        })


# ============================================================
# 按英雄编号排序
# ============================================================

heroes.sort(
    key=lambda x: x["hero_id"]
)


# ============================================================
# 显示统计
# ============================================================

print("=" * 70)

print(
    f"王者荣耀英雄原画查看器："
    f"Excel 共 {total_hero_count} 个英雄，"
    f"{len(heroes)} 个英雄有待处理的海报"
)

print("=" * 70)


# ============================================================
# 没有需要处理的英雄：直接结束
# （GUI 里 pause 会跳过，输出完立即退出）
# ============================================================

if not heroes:

    print(
        "没有需要处理的英雄，"
        "所有海报都已填写皮肤名。"
    )

    pause()

    exit()


# ============================================================
# 一个英雄一个英雄处理
# ============================================================

# 是否被 Ctrl+C 手动中断
interrupted = False


try:

    for index, hero in enumerate(
        heroes
    ):

        hero_id = hero["hero_id"]

        hero_name = hero["hero_name"]

        skins = hero["skins"]


        print(
            f"[{index + 1}/{len(heroes)}] "
            f"{hero_id:03d} {hero_name}  "
            "待查看格子："
            + " ".join(
                f"{skin['skin_id']:02d}"
                for skin in skins
            )
        )


        # ========================================================
        # 直接打开当前英雄未处理的海报
        # ========================================================

        for skin in skins:

            print(
                f"{skin['skin_id']:02d} → "
                f"{skin['url']}"
            )

            webbrowser.open_new_tab(
                skin["url"]
            )


        print(
            f"{hero_name} 的 "
            f"{len(skins)} 张未处理海报已打开。"
        )


        # ========================================================
        # 等待用户处理完当前英雄
        #
        # 提示用 print 先发出来：input 的提示词走管道
        # 不会实时显示，还会和下一段输出粘在同一行，
        # 所以 input() 本身不带提示词
        # ========================================================

        if index < len(heroes) - 1:

            print(
                "在 Excel 里填好皮肤名，"
                "然后按回车（GUI 里点「继续」）"
                "进入下一个英雄。"
            )

            input()


except KeyboardInterrupt:

    interrupted = True

    print(
        "已手动中断，"
        "已填写的内容保留在 Excel 里。"
    )


# ============================================================
# 完成
# ============================================================

if interrupted:

    print(
        "本次处理被中断，"
        "下次运行会从还没填写的格子继续。"
    )

else:

    print(
        f"全部 {len(heroes)} 个英雄的"
        f"海报已查看完成。"
    )

pause()